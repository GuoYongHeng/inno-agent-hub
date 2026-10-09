"""Native editable and image-mode presentation rendering and asset preflight."""
from pathlib import Path


def prepare(content, input_dir):
    from PIL import Image
    mode = content.get('mode', 'editable')
    seen = set()
    for slide in content['slides']:
        if mode == 'image' and slide['type'] != 'image':
            raise ValueError('image mode requires image slides')
        if mode == 'editable' and slide['type'] == 'image':
            raise ValueError('full-page image slides require image mode')
        if slide['type'] == 'table':
            lengths = {len(row) for row in slide['table']}
            if len(lengths) != 1:
                raise ValueError('table rows must have equal column counts')
        if 'image' in slide:
            path = (input_dir / slide['image']).resolve()
            if mode == 'image' and path in seen:
                raise ValueError('duplicate slide image')
            seen.add(path)
            with Image.open(path) as img:
                width, height = img.size
                img.verify()
            if mode == 'image' and abs(width / height - 16 / 9) > .01:
                raise ValueError('image slide must be 16:9')
            slide['image'] = str(path)


def render(data, output):
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = data['content']['title']
    for spec in data['content']['slides']:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if spec['type'] == 'image':
            slide.shapes.add_picture(spec['image'], 0, 0, prs.slide_width, prs.slide_height)
        else:
            title = slide.shapes.add_textbox(Inches(.6), Inches(.3), Inches(12.1), Inches(.9))
            title.text_frame.word_wrap = True
            title.text_frame.text = spec['title']
            for run in title.text_frame.paragraphs[0].runs:
                run.font.size, run.font.bold, run.font.name = Pt(28), True, 'PingFang SC'
            width = 7.2 if spec.get('image') else 11.9
            if spec['type'] == 'text':
                box = slide.shapes.add_textbox(Inches(.7), Inches(1.4), Inches(width), Inches(5.4))
                tf = box.text_frame
                tf.word_wrap = True
                for i, line in enumerate(spec['body']):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.size, p.font.name = Pt(24), 'PingFang SC'
                    p.space_after = Pt(16)
            elif spec['type'] == 'table':
                rows = spec['table']
                table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.7), Inches(1.4), Inches(width), Inches(5)).table
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        cell = table.cell(r, c)
                        cell.text = val
                        for p in cell.text_frame.paragraphs:
                            p.font.size, p.font.name = Pt(20), 'PingFang SC'
            elif spec['type'] == 'chart':
                cd = ChartData()
                cd.categories = spec['chart']['categories']
                for series in spec['chart']['series']:
                    cd.add_series(series['name'], series['values'])
                slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.4), Inches(width), Inches(5), cd)
            if spec.get('image'):
                from PIL import Image
                with Image.open(spec['image']) as img:
                    ratio = img.width / img.height
                w, h = min(4.1, 5 / (1 / ratio)), min(5, 4.1 / ratio)
                slide.shapes.add_picture(spec['image'], Inches(8.6), Inches(1.4), Inches(w), Inches(h))
        slide.notes_slide.notes_text_frame.text = f"slide_id:{spec['id']}\n{spec.get('notes', '')}"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
