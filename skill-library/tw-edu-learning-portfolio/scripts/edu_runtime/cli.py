from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
import re
import sys
import uuid
from pathlib import Path
from typing import Any

SKILLS = {
    "tw-edu-lesson-plan-108": "docx", "tw-edu-curriculum-mapper": "xlsx",
    "tw-edu-exam-generator": "exam", "tw-edu-rubric-designer": "docx",
    "tw-edu-feedback-writer": "docx", "tw-edu-learning-portfolio": "docx",
    "tw-edu-classroom-culture": "docx", "tw-edu-differentiated": "docx",
    "tw-edu-formative-assessment": "docx", "tw-edu-interdisciplinary": "docx",
    "tw-edu-meeting-facilitator": "docx", "tw-edu-parent-communication": "docx",
    "tw-edu-pbl-designer": "docx", "tw-edu-school-document": "docx",
    "tw-edu-worksheet-creator": "docx", "tw-edu-anti-ai-assessment": "docx",
    "tw-edu-mini-app": "html", "tw-edu-research-viz": "png",
    "tw-edu-slides-creator": "pptx",
}

class InputError(Exception): pass

def _json(path: Path) -> Any:
    try: return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e: raise InputError(f"cannot read JSON {path}: {e}") from e

def _hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(65536), b""): h.update(block)
    return h.hexdigest()

def _skill_dir() -> Path:
    # <skill>/scripts/edu_runtime/cli.py
    return Path(__file__).resolve().parents[2]

def _validate(data: dict, schema_path: Path) -> None:
    def check_values(value, location='$'):
        if isinstance(value, str) and not value.strip():
            raise InputError(f'empty text at {location}')
        if isinstance(value, float) and not math.isfinite(value):
            raise InputError(f'non-finite number at {location}')
        if isinstance(value, dict):
            for key, item in value.items(): check_values(item, f'{location}.{key}')
        elif isinstance(value, list):
            for index, item in enumerate(value): check_values(item, f'{location}[{index}]')
    check_values(data)
    schema = _json(schema_path)
    try:
        import jsonschema
    except ImportError:
        from .schema_fallback import SchemaError, validate
        try: validate(data, schema)
        except SchemaError as e: raise InputError(f"validation failed: {e}") from e
    else:
        try: jsonschema.validate(data, schema)
        except jsonschema.ValidationError as e:
            location = ".".join(str(p) for p in e.absolute_path) or "$"
            raise InputError(f"validation failed at {location}: {e.message}") from e

def _semantic(skill: str, d: dict) -> None:
    c = d["content"]
    def unique(items, field, label):
        values=[x[field] for x in items]
        if len(values)!=len(set(values)): raise InputError(f"duplicate {label}")
    if skill == "tw-edu-lesson-plan-108":
        unique(c["objectives"],"id","objective id"); unique(c["activities"],"id","activity id"); unique(c["assessments"],"id","assessment id")
        ids = {x["id"] for x in c["objectives"]}
        if sum(x["minutes"] for x in c["activities"]) != c["total_minutes"]: raise InputError("activity minutes must equal total_minutes")
        for group in (c["activities"], c["assessments"]):
            if any(not set(x["objective_ids"]).issubset(ids) for x in group): raise InputError("unknown objective_id")
        taught=set().union(*(set(x["objective_ids"]) for x in c["activities"])); assessed=set().union(*(set(x["objective_ids"]) for x in c["assessments"]))
        if taught != ids: raise InputError(f"every objective must be taught; uncovered: {sorted(ids-taught)}")
        if assessed != ids: raise InputError(f"every objective must be assessed; uncovered: {sorted(ids-assessed)}")
        for code in c.get("curriculum_codes", []):
            if not code.get("verified", False) and not code.get("verification_note"): raise InputError("unverified curriculum code requires verification_note")
    elif skill == "tw-edu-exam-generator":
        unique(c["questions"],"id","question id")
        if len(c["questions"]) != c["expected_question_count"]: raise InputError("actual question count mismatch")
        if sum(q["points"] for q in c["questions"]) != c["total_points"]: raise InputError("question points do not equal total_points")
        for q in c["questions"]:
            if q.get("options"): unique(q["options"],"id",f"option id in question {q['id']}")
            if q["type"] == "multiple_choice" and (q.get("answer") not in [o["id"] for o in q.get("options", [])]): raise InputError(f"question {q['id']} answer is not an option")
    elif skill == "tw-edu-rubric-designer":
        levels = c["levels"]
        unique(levels, 'id', 'level id')
        if c["type"] == "analytic":
            unique(c['dimensions'], 'id', 'dimension id')
            if any(set(x["descriptions"]) != {l["id"] for l in levels} for x in c["dimensions"]): raise InputError("each analytic dimension needs every level description")
            if sum(x["weight"] for x in c["dimensions"]) != c["total_points"]: raise InputError("dimension weights do not equal total_points")
        elif set(c["descriptions"]) != {l["id"] for l in levels}: raise InputError("holistic rubric needs every level description")
        elif max(l["score"] for l in levels) != c["total_points"]: raise InputError("holistic total_points must equal the highest level score")
    elif skill == "tw-edu-curriculum-mapper":
        for unit in c["units"]:
            for code in unit["codes"]:
                if not code["verified"] and not code.get("verification_note"): raise InputError("unverified curriculum code requires verification_note")
    elif skill == "tw-edu-research-viz" and c["type"] == "prisma":
        p = c["prisma"]
        if p["identified"] != p["duplicates_removed"] + p["screened"]: raise InputError("PRISMA identification counts are not conserved")
        if p["screened"] != p["screening_excluded"] + p["full_text_assessed"]: raise InputError("PRISMA screening counts are not conserved")
        if p["full_text_assessed"] != p["full_text_excluded"] + p["included"]: raise InputError("PRISMA eligibility counts are not conserved")
    elif skill == "tw-edu-mini-app" and c["mode"] == "quiz":
        unique(c["questions"],"id","question id")
        for q in c["questions"]:
            unique(q['options'], 'id', 'option id')
            if q["answer"] not in [o["id"] for o in q["options"]]: raise InputError(f"question {q['id']} answer is not an option")
    elif skill == "tw-edu-anti-ai-assessment":
        for item in c["items"]:
            if sum(x["score"] for x in item["dimensions"]) != item["total_score"]: raise InputError(f"item {item['id']} dimension scores do not equal total_score")
    elif skill == "tw-edu-slides-creator":
        ids=[x["id"] for x in c["slides"]]
        if ids != list(range(1,len(ids)+1)): raise InputError("slide ids must be unique and contiguous from 1")
        for slide in c["slides"]:
            if slide["type"]=="chart":
                n=len(slide["chart"]["categories"])
                if any(len(s["values"])!=n for s in slide["chart"]["series"]): raise InputError(f"slide {slide['id']} chart category/value lengths differ")

DOC_LABELS={
"tw-edu-feedback-writer":{"students":"学生观察与反馈"},"tw-edu-learning-portfolio":{"records":"学习历程纪录"},
"tw-edu-classroom-culture":{"agreements":"班级共识","routines":"日常程序","response_plan":"事件回应计划"},
"tw-edu-differentiated":{"shared_goal":"共同学习目标","learner_groups":"学习需求与支持","activities":"差异化活动"},
"tw-edu-formative-assessment":{"learning_target":"学习目标","checks":"评价检核","response_rules":"依证据调整教学"},
"tw-edu-interdisciplinary":{"disciplines":"跨域学科","driving_question":"驱动问题","discipline_contributions":"各科贡献","activities":"学习活动","product":"成果作品"},
"tw-edu-meeting-facilitator":{"participants":"与会人员","agenda":"议程","decisions":"决议","actions":"待办追踪"},
"tw-edu-parent-communication":{"recipients":"收件对象","purpose":"沟通目的","message":"讯息内容","requested_action":"期待配合事项","contact_channel":"联络管道"},
"tw-edu-pbl-designer":{"driving_question":"驱动问题","authentic_context":"真实情境","milestones":"里程碑","final_product":"最终成果","assessment_criteria":"评价准则"},
"tw-edu-school-document":{"document_type":"文件类型","basis":"依据","purpose":"目的","implementation":"实施方式","responsible_people":"权责人员","expected_results":"预期成果"},
"tw-edu-worksheet-creator":{"instructions":"作答说明","prompts":"学习任务","reflection":"反思问题"},
"tw-edu-anti-ai-assessment":{"items":"评价项目、维度分数与设计理由"},
}
FIELD_LABELS={"id":"编号","title":"名称","instructions":"进行方式","materials":"材料","date":"日期","artifact":"作品","evidence":"证据","reflection":"反思","next_step":"下一步","situation":"情境","steps":"步骤","trigger":"触发情况","teacher_action":"教师行动","follow_up":"后续追踪","support":"支持方式","prompt":"题目","success_criteria":"成功准则","evidence_capture":"证据搜集","action":"行动","discipline":"学科","knowledge":"知识内容","method":"方法","topic":"议题","minutes":"分钟","owner":"负责人","due":"期限","deliverable":"交付成果","feedback":"反馈方式","item":"项目","details":"内容","response_space_lines":"作答行数","student_id":"学生代码","observations":"具体观察","strengths":"优势","next_steps":"下一步","assessment_item":"评价题目","dimensions":"评分维度","total_score":"总分","reason":"理由","redesign":"调整方案","name":"名称","score":"分数"}

def _strings(value: Any, prefix=""):
    if isinstance(value, dict):
        for k, v in value.items(): yield from _strings(v, f"{prefix}{FIELD_LABELS.get(k, k)}｜")
    elif isinstance(value, list):
        for i, v in enumerate(value, 1): yield from _strings(v, f"{prefix}{i}｜")
    elif value is not None: yield prefix.rstrip("｜"), str(value)

def _docx(d: dict, output: Path, teacher=False) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document(); sec = doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=sec.bottom_margin=Cm(1.8); sec.left_margin=sec.right_margin=Cm(2)
    title = d["content"].get("title") or d["context"].get("topic") or d["skill"]
    doc.add_heading(title + ("｜教师答案" if teacher else ""), 0)
    doc.add_paragraph(f"{d['context'].get('subject','')}　{d['context'].get('grade','')}")
    doc.add_heading("来源与查核状态",1)
    for source in d["sources"]: doc.add_paragraph(f"{source['title']}｜{source['status']}"+(f"｜{source['url']}" if source.get('url') else ""))
    if d["skill"] == "tw-edu-lesson-plan-108":
        c=d["content"]
        for heading,items,columns in [("学习目标",c["objectives"],["id","description"]),("教学活动",c["activities"],["id","title","minutes","instructions","objective_ids"]),("评价设计",c["assessments"],["id","method","criteria","objective_ids"])]:
            doc.add_heading(heading,1); table=doc.add_table(rows=1,cols=len(columns)); table.style="Table Grid"
            for j,key in enumerate(columns): table.rows[0].cells[j].text={"id":"编号","description":"目标叙述","title":"活动","minutes":"分钟","instructions":"进行方式","objective_ids":"对应目标","method":"方法","criteria":"成功准则"}[key]
            for item in items:
                cells=table.add_row().cells
                for j,key in enumerate(columns): cells[j].text="、".join(item[key]) if isinstance(item[key],list) else str(item[key])
        if c.get("curriculum_codes"):
            doc.add_heading("课程标准代码与查核",1)
            for code in c["curriculum_codes"]: doc.add_paragraph(f"{code['code']}｜{'已查核' if code['verified'] else '待查核'}｜{code['description']}｜{code.get('verification_note','')}")
    elif d["skill"] == "tw-edu-rubric-designer":
        c=d["content"]; doc.add_heading("评价量规",1); levels=c["levels"]
        if c["type"]=="analytic":
            table=doc.add_table(rows=1,cols=2+len(levels)); table.style="Table Grid"
            heads=["维度","配分"]+[f"{x['label']}（{x['score']}）" for x in levels]
            for i,x in enumerate(heads): table.rows[0].cells[i].text=x
            for dim in c["dimensions"]:
                cells=table.add_row().cells; cells[0].text=dim["name"]; cells[1].text=str(dim["weight"])
                for i,lvl in enumerate(levels,2): cells[i].text=dim["descriptions"][lvl["id"]]
        else:
            table=doc.add_table(rows=1,cols=3); table.style="Table Grid"
            for i,x in enumerate(["等第","分数","整体描述"]): table.rows[0].cells[i].text=x
            for lvl in levels:
                cells=table.add_row().cells; cells[0].text=lvl["label"]; cells[1].text=str(lvl["score"]); cells[2].text=c["descriptions"][lvl["id"]]
        doc.add_paragraph(f"总分：{c['total_points']}")
    elif d["skill"] == "tw-edu-feedback-writer":
        for student in d["content"]["students"]:
            doc.add_heading(f"学生纪录：{student['student_id']}",1)
            table=doc.add_table(rows=4,cols=2); table.style="Table Grid"
            for row,(label,key) in zip(table.rows,[("具体观察","observations"),("优势","strengths"),("下一步","next_steps"),("反馈文字","feedback")]): row.cells[0].text=label; row.cells[1].text="\n".join(student[key]) if isinstance(student[key],list) else student[key]
    elif d["skill"] == "tw-edu-exam-generator":
        c=d["content"]
        doc.add_paragraph(f"题数：{len(c['questions'])} 题　总分：{c['total_points']} 分")
        for i,q in enumerate(c["questions"],1):
            doc.add_paragraph({'multiple_choice':'选择题','short_answer':'简答题','essay':'申论题','true_false':'是非题'}[q['type']])
            doc.add_heading(f"{i}. {q['prompt']}（{q['points']} 分）", 2)
            for o in q.get("options",[]): doc.add_paragraph(f"{o['id']}. {o['text']}")
            if teacher:
                doc.add_paragraph(f"答案：{q['answer']}")
                if q.get("explanation"): doc.add_paragraph(f"解析：{q['explanation']}")
            else: doc.add_paragraph("作答：________________________________")
    else:
        for key,val in d["content"].items():
            if key=="title": continue
            doc.add_heading(DOC_LABELS.get(d["skill"],{}).get(key,str(key)), 1)
            if isinstance(val,list):
                for item in val:
                    if isinstance(item,dict):
                        p=doc.add_paragraph(style="List Bullet")
                        p.add_run("；".join(f"{FIELD_LABELS.get(k,k)}：{v}" for k,v in item.items() if not isinstance(v,(list,dict))))
                        for k,v in item.items():
                            if isinstance(v,(list,dict)):
                                for label,text in _strings(v,f"{FIELD_LABELS.get(k,k)}｜"): doc.add_paragraph(f"{label}：{text}",style="List Bullet 2")
                        if d["skill"]=="tw-edu-worksheet-creator":
                            for _ in range(item.get("response_space_lines",0)): doc.add_paragraph("________________________________________________")
                    else: doc.add_paragraph(str(item),style="List Bullet")
            elif isinstance(val,dict):
                for label,text in _strings(val): doc.add_paragraph(f"{label}：{text}")
            else: doc.add_paragraph(str(val))
    styles=doc.styles
    for n in ["Normal","Title","Heading 1","Heading 2"]:
        styles[n].font.name="PingFang SC"; styles[n]._element.rPr.rFonts.set(qn("w:eastAsia"),"PingFang SC"); styles[n].font.size=Pt(11 if n=="Normal" else 16)
        styles[n].paragraph_format.line_spacing=1.35
        styles[n].font.color.rgb = RGBColor(0, 0, 0)
        for border in styles[n]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in paragraph._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for table in doc.tables:
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        headings = [cell.text for cell in table.rows[0].cells]
        weights = [1 if h in {'编号','分钟','配分'} else 2 if h == '对应目标' else 4 for h in headings]
        table.autofit = False
        for j, weight in enumerate(weights):
            table.columns[j].width = Cm(17 * weight / sum(weights))
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(17 * weights[j] / sum(weights))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = 'PingFang SC'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name="PingFang SC"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"PingFang SC")
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)

def _xlsx(d: dict, output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    c=d["content"]; wb=Workbook(); ws=wb.active; ws.title="课程地图"
    ws.append([c["title"]]); ws.merge_cells("A1:F1"); ws["A1"].font=Font(bold=True,size=16)
    ws.append(["科目",d["context"]["subject"],"年级",d["context"]["grade"],"主题",d["context"]["topic"]])
    ws.append(["来源", "；".join(f"{x['title']}({x['status']})" for x in d["sources"])])
    headers=["单元","周次","节数","目标","课程标准代码","评价"]
    ws.append(headers)
    for cell in ws[4]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1A5276")
    for u in c["units"]: ws.append([u["name"],u["weeks"],u["periods"],"\n".join(u["goals"]),"\n".join(f"{x['code']}｜{'已查核' if x['verified'] else '待查核'}｜{x.get('verification_note','')}" for x in u["codes"]),"\n".join(u["assessments"])])
    for col,w in zip("ABCDEF",[22,12,8,40,22,35]): ws.column_dimensions[col].width=w
    for row in ws.iter_rows():
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes = 'A5'
    ws.print_title_rows = '1:4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)

def _safe_json(v): return json.dumps(v,ensure_ascii=False).replace("<","\\u003c")

def _html(d: dict, output: Path) -> None:
    c=d["content"]; mode=c["mode"]; payload=_safe_json(c)
    script = Path(__file__).with_name('miniapp.js').read_text(encoding='utf-8')
    title=html.escape(c["title"])
    page=f'''<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{title}</title><style>body{{font-family:system-ui,sans-serif;max-width:56rem;margin:auto;padding:2rem;background:#f7fafc;color:#17202a}}button{{display:block;margin:.6rem 0;padding:.8rem 1rem}}button:focus{{outline:3px solid #2471a3}}</style><h1>{title}</h1><p>{html.escape(d['context']['subject'])}｜{html.escape(d['context']['grade'])}</p><main id="app"></main><script id="data" type="application/json">{payload}</script><script>{script}</script></html>'''
    output.parent.mkdir(parents=True,exist_ok=True); output.write_text(page,encoding="utf-8")

def _font() -> str:
    roots=[Path("/System/Library/Fonts"),Path("/Library/Fonts"),Path("/usr/share/fonts"),Path.home()/"Library/Fonts"]
    preferred=("NotoSansCJK","NotoSansTC","PingFang","JhengHei","Songti")
    for root in roots:
        if root.exists():
            files=[p for p in root.rglob("*") if p.suffix.lower() in {".ttf",".otf",".ttc"}]
            for token in preferred:
                for path in files:
                    if token.lower() in path.name.lower(): return str(path)
    raise InputError("CJK font unavailable; install Noto Sans CJK or another Traditional Chinese font")

def _png(d: dict, output: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont
    c=d["content"]; font_path=_font(); title_font=ImageFont.truetype(font_path,34); body_font=ImageFont.truetype(font_path,24)
    if c["type"]=="prisma":
        p=c["prisma"]; labels=[("辨识（去除重复 "+str(p["duplicates_removed"])+"）",p["identified"]),("筛选（排除 "+str(p["screening_excluded"])+"）",p["screened"]),("全文审查（排除 "+str(p["full_text_excluded"])+"）",p["full_text_assessed"]),("纳入",p["included"])]
    else: labels=[(x["label"],x.get("value","")) for x in c["nodes"]]
    image=Image.new("RGB",(1600,900),"white"); draw=ImageDraw.Draw(image)
    draw.text((800,45),c["title"]+"（简易流程图）",font=title_font,fill="#1A5276",anchor="ma")
    gap=700/max(1,len(labels)); box_h=min(105,gap*.7)
    for i,(label,value) in enumerate(labels):
        y=125+i*gap; draw.rounded_rectangle((480,y,1120,y+box_h),radius=16,fill="#EBF5FB",outline="#2471A3",width=3)
        draw.multiline_text((800,y+box_h/2),f"{label}\n{value}",font=body_font,fill="#17202A",anchor="mm",align="center")
        if i<len(labels)-1: draw.line((800,y+box_h,800,y+gap),fill="#2471A3",width=4); draw.polygon([(790,y+gap-12),(810,y+gap-12),(800,y+gap)],fill="#2471A3")
    output.parent.mkdir(parents=True,exist_ok=True); image.save(output)

def _report(input_path: Path, outputs: list[Path], d: dict, sample: bool) -> Path:
    primary=outputs[0]; report=primary.with_name(primary.name+".validation.json")
    checks=[]
    if any(x["status"]!="verified" for x in d["sources"]): checks.append("source_verification")
    if d["skill"]=="tw-edu-lesson-plan-108" and any(not x.get("verified") for x in d["content"].get("curriculum_codes",[])): checks.append("curriculum_code_verification")
    if d["skill"]=="tw-edu-curriculum-mapper" and any(not x.get("verified") for u in d["content"]["units"] for x in u["codes"]): checks.append("curriculum_code_verification")
    checks.append("human_visual_review")
    body={"schema_version":"1.0","skill":d["skill"],"sample":sample,"input":{"path":str(input_path),"sha256":_hash(input_path)},"outputs":[{"path":str(x),"sha256":_hash(x)} for x in outputs],"pending_checks":checks}
    report.write_text(json.dumps(body,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); return report

def run(skill: str, input_path: Path, output: Path, validate_only=False, sample=False) -> list[Path]:
    root=_skill_dir(); schema=root/"schemas"/"input.schema.json"; d=_json(input_path)
    _validate(d,schema)
    if d.get("skill") != skill: raise InputError(f"input skill must be {skill}")
    _semantic(skill,d)
    if skill == 'tw-edu-slides-creator':
        from .slides import prepare
        try: prepare(d['content'], input_path.resolve().parent)
        except (ValueError, OSError) as exc: raise InputError(str(exc)) from exc
    if validate_only: return []
    if sample and '范例' not in d['content']['title']:
        d['content']['title'] += '（范例）'
    if output.exists(): raise InputError(f'output already exists: {output}')
    kind=SKILLS[skill]; outputs=[]
    if kind=="exam":
        stem=output.with_suffix(""); student=stem.with_name(stem.name+"-student").with_suffix(".docx"); teacher=stem.with_name(stem.name+"-teacher").with_suffix(".docx")
        if student.exists() or teacher.exists(): raise InputError('exam output already exists')
        _docx(d,student,False); _docx(d,teacher,True); outputs=[student,teacher]
    else:
        expected={"docx":".docx","xlsx":".xlsx","html":".html","png":".png","pptx":".pptx"}[kind]
        if output.suffix.lower()!=expected: raise InputError(f"output must use {expected}")
        from .slides import render as render_slides
        {"docx":_docx,"xlsx":_xlsx,"html":_html,"png":_png,"pptx":render_slides}[kind](d,output); outputs=[output]
    _report(input_path,outputs,d,sample); return outputs

def main(skill_name: str) -> int:
    if skill_name not in SKILLS: print(f"unsupported skill: {skill_name}",file=sys.stderr); return 2
    p=argparse.ArgumentParser(description=f"Input-driven generator for {skill_name}")
    p.add_argument("--input",type=Path); p.add_argument("--output",type=Path); p.add_argument("--validate-only",action="store_true"); p.add_argument("--example",action="store_true")
    args,unknown=p.parse_known_args()
    if unknown: p.error("legacy/unknown flags are unsupported; migrate to --input, --output, --validate-only, or --example")
    if args.example and args.input: p.error("--example and --input are mutually exclusive")
    if not args.example and not args.input: p.error("--input is required unless --example is used")
    if not args.validate_only and not args.output:
        extension = 'docx' if SKILLS[skill_name] == 'exam' else SKILLS[skill_name]
        args.output = Path.cwd() / 'artifacts' / skill_name / uuid.uuid4().hex[:12] / ('output.' + extension)
    source=_skill_dir()/"examples"/"example.json" if args.example else args.input
    try:
        outputs=run(skill_name,source,args.output,args.validate_only,args.example)
        print("valid" if args.validate_only else "generated: "+", ".join(str(x) for x in outputs)); return 0
    except (InputError, OSError, ValueError) as e: print(f"error: {e}",file=sys.stderr); return 2
