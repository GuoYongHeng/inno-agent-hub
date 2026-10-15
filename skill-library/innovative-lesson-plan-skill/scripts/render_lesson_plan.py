"""创新教案生成主入口。

整条链路：
    校验请求 -> 拼 prompt -> 调用生成模型 -> 抽取 JSON -> 结构修补
    -> 渲染 Word-ready Markdown -> 导出 DOCX -> 规则校验

四个值得说明的设计决定：

1. 渲染由本地脚本完成，而不是让模型直接写 Markdown。模型只负责教学内容，
   版式交给代码，这样三类教案的章节顺序、表格列名不会每次都不一样。
2. 支持 --from-json。没有 API key、或者只想改版式时，可以直接对着一份已有的
   教案 JSON 重跑渲染与校验。这让整条链路能在离线环境里自检。
3. 交付契约：一次生成只交付一个 Word 文件，文件名由教案信息自动派生为中文
   （如「六年级科学《校园节水方案设计》项目化学习教案.docx」）。JSON 与 Markdown
   仍然会生成，但它们属于校验与渲染的中间产物，写进输出目录下的「中间产物」子目录。
4. 课标与教材内容由生成模型内置的学科知识直接给出，不依赖任何外部课标库，
   也不在教案里留「请教师核对课标」这类把活推回给教师的提示。

用法：

    python3 scripts/render_lesson_plan.py <输出目录> --config <请求 JSON>
    python3 scripts/render_lesson_plan.py <输出目录> --sample pbl
    python3 scripts/render_lesson_plan.py <输出目录> --from-json <教案 JSON> --request <请求 JSON>
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lp_common import (  # noqa: E402
    EXAMPLES_DIR,
    INNOVATION_TYPES,
    INTERMEDIATE_DIRNAME,
    TYPE_DESIGN_FIELDS,
    TYPE_LABELS,
    compact_json,
    deliverable_file_name,
    ensure_list,
    item_get,
    load_env_file,
    load_json,
    load_reference_data,
    non_empty,
    normalize_innovation_type,
    sanitize_file_name,
    table_cell,
    value_to_text,
)
from normalize_lesson_plan import normalize_lesson_plan  # noqa: E402
from validate_lesson_plan import validate_document  # noqa: E402

load_env_file()

BASE_URL = (
    os.getenv("LESSON_PLAN_BASE_URL")
    or os.getenv("DASHSCOPE_BASE_URL")
    or os.getenv("QWEN_BASE_URL")
    or "https://dashscope.aliyuncs.com/compatible-mode/v1"
)
MODEL = os.getenv("LESSON_PLAN_MODEL") or os.getenv("GENERATOR_MODEL") or os.getenv("QWEN_MODEL") or "qwen3.5-122b-a10b"
DEBUG_DIR = Path(os.getenv("LESSON_PLAN_DEBUG_DIR", "/tmp/lesson_plan_debug"))

# 每个样例都指向真实存在的文件。指向不存在的文件是「文档承诺了却跑不通」的典型，
# 与其让老师拿到报错，不如让这里保持可验证。
SAMPLE_CONFIGS = {
    "pbl": EXAMPLES_DIR / "pbl-campus-water-saving.json",
    "interdisciplinary": EXAMPLES_DIR / "interdisciplinary-campus-water-data.json",
    "ai": EXAMPLES_DIR / "ai-integrated-expository-writing.json",
}


OUTPUT_CONTRACT: Dict[str, Any] = {
    "lessonMeta": {
        "subject": "请求中的学科；跨学科可写成「语文 + 化学」",
        "grade": "请求中的年级",
        "topic": "请求中的课题或真实任务主题",
        "innovationType": "PBL | interdisciplinary | ai_integrated",
        "durationMin": 40,
        "textbookVersion": "教师确认的教材版本（取自 references/textbook-versions.json，不得自行假设）",
        "periodNo": "第几课时，如「第1课时」",
        "linkedSubject": "跨学科必填：关联学科名称",
        "lessonType": "课型的中文表述，如「项目化学习课」「跨学科融合课」",
    },
    "backgroundAnalysis": {
        "_doc": "按类型填写：PBL 填 textbookPosition/priorKnowledge/inquiryExperience/collaborationAbility；跨学科填 primarySubject 与 linkedSubject（对象，含 unitPosition 与 standardRequirement）以及 curriculumIntersection；AI 融合填 textbookPosition 与 aiInterventionRationale",
    },
    "studentAnalysis": {
        "_doc": "PBL 不使用本字段（学情合入 backgroundAnalysis）；跨学科填 primarySubjectReadiness/linkedSubjectReadiness/crossSubjectExperience；AI 融合填 priorKnowledge/aiToolExperience/critiqueAbility/independentThinkingHabit",
    },
    "coreCompetencies": [
        {
            "id": "cc-1",
            "dimension": "课程标准的核心素养维度（如：科学观念、科学思维、探究实践、态度责任）",
            "target": "本课要落实的具体表现",
        }
    ],
    "objectives": [
        {
            "id": "obj-1",
            "description": "使用可观察行为动词表述的学习目标",
            "behaviorVerb": "取自 action-verbs.json 的 allowedVerbs",
            "owningSubject": "跨学科必填：本目标归属的学科，取值为「语文」「化学」或「语文+化学」",
            "linkedActivities": ["act-1"],
            "assessmentEvidence": ["rubric-1"],
        }
    ],
    "teachingFocus": "教学重点，必须对齐主目标（PBL 是驱动问题的探究过程；跨学科是融合节点；AI 融合是 AI 服务的目标达成与审辨）",
    "teachingDifficulty": "教学难点，必须与学情描述指向同一处困难，并在流程中有支架",
    "innovationDesign": {
        "type": "与 innovationType 一致",
        "_doc": "按类型填写：PBL 必填 drivingQuestion/milestones/finalProduct；跨学科必填 disciplineConnections/integrationNodes/commonProduct/applicableBoundary；AI 融合必填 aiToolRoles/interventionStages/useBoundaries/studentCritiqueTasks。milestones、disciplineConnections、integrationNodes、aiToolRoles、interventionStages、studentCritiqueTasks 必须是数组",
    },
    "activityFlow": [
        {
            "id": "act-1",
            "stage": "环节或里程碑名称",
            "durationMin": 5,
            "subjectTag": "跨学科必填：本环节的学科归属，取值为「语文」「化学」或「学科交汇」",
            "executor": "AI 融合必填：本环节的执行者，取值为「教师」「学生」「AI」或「人机协同」",
            "teacherActions": ["教师活动"],
            "studentActions": ["学生活动"],
            "designIntent": "本环节为什么必须这样安排、为什么不能删",
            "outputs": ["学生或小组产出"],
            "assessmentLinks": ["rubric-1"],
        }
    ],
    "assessmentRubric": [
        {
            "id": "rubric-1",
            "category": "PBL 分过程评价与作品评价时填 process 或 product，其余类型可省略",
            "dimension": "评价维度",
            "excellent": "优秀表现",
            "qualified": "合格表现",
            "needsImprovement": "待改进表现",
            "evidence": "证据来源",
        }
    ],
    "resources": [{"type": "资源类型", "name": "资源名称", "usage": "使用方式"}],
    "export": {"format": "markdown", "markdown": "留空，脚本会重建"},
    "qualityReport": {"checks": [], "warnings": []},
}


SYSTEM_PROMPT = """你是一位熟悉中国中小学课堂的一线教研员，正在为教师撰写可直接使用的创新教案。

你写教案时有三个坚持：
- 创新点必须服务学习目标：没有真实驱动问题、真实融合节点或真实审辨任务的课，不算创新课；
- 目标必须可观察，不用「了解、掌握、培养、提高」这类无法检验的动词；
- 每个环节都要有学生真正在做事，并写清设计意图，不能只记录教师做了什么。

你只输出一个 JSON 对象，不输出解释、不输出多个候选方案。"""


def validate_request_config(request: Dict[str, Any]) -> Tuple[List[str], List[str]]:
    """请求预检。缺必填字段要拦下来，缺推荐字段只提示。"""
    failures: List[str] = []
    warnings: List[str] = []

    if not isinstance(request, dict):
        return ["请求 JSON 顶层必须是对象"], []

    for key in ("subject", "grade", "topic"):
        if not non_empty(request.get(key)):
            failures.append(f"缺少必填字段: {key}")

    raw_type = request.get("innovationType")
    if not non_empty(raw_type):
        failures.append("缺少必填字段: innovationType（应先由 Controller 完成创新类型澄清）")
    else:
        innovation_type = normalize_innovation_type(raw_type)
        if not innovation_type:
            failures.append(
                f"innovationType 无法识别: {raw_type}（应为 {'/'.join(INNOVATION_TYPES)} 或对应中文名称）"
            )
        elif innovation_type != raw_type:
            request["innovationType"] = innovation_type

    duration = request.get("durationMin")
    if duration is None:
        request["durationMin"] = 40
        warnings.append("未提供 durationMin，已按 40 分钟生成并写入 assumptions")
    else:
        try:
            if float(duration) <= 0:
                failures.append("durationMin 必须大于 0")
        except (TypeError, ValueError):
            failures.append("durationMin 必须是数字")

    if not non_empty(request.get("constraints")):
        warnings.append("未提供 constraints，将按普通教室条件保守生成")
    if not non_empty(request.get("requirements")):
        warnings.append("未提供 requirements，将只按创新类型默认闭环生成")
    if not non_empty(request.get("textbookVersion")):
        # 教材版本必须先问后生成：即使该学科在清单里只有一个版本，
        # 也要向教师提问确认，「只有一个版本」不能当成不问的理由。
        failures.append(
            "缺少必填字段: textbookVersion。必须先向教师确认教材版本"
            "（按 references/textbook-versions.json 提问，即使该学科只有一个可选版本也要确认），"
            "不得默认或猜测某个版本"
        )

    context = request.get("confirmedContext")
    if context is None:
        warnings.append("未提供 confirmedContext，教学设计的个性化程度会降低")
    elif not isinstance(context, dict):
        failures.append("confirmedContext 必须是对象")
    else:
        from lp_common import TYPE_CONTEXT_FIELDS

        lesson_type = normalize_innovation_type(request.get("innovationType"))
        missing = [field for field in TYPE_CONTEXT_FIELDS.get(lesson_type, []) if not non_empty(context.get(field))]
        if missing:
            warnings.append(f"confirmedContext 缺少类型专属字段: {', '.join(missing)}；将使用预设默认值")

    return failures, warnings


def _type_requirements(innovation_type: str) -> str:
    return {
        "PBL": "驱动问题必须是一个学生需要探究才能回答的开放问题（不能靠搜索直接抄到答案）；里程碑要覆盖问题提出、探究、产出、评价四段，并说明每个里程碑的阶段产出。",
        "interdisciplinary": "必须说明两学科在课标层面的交汇依据、融合节点、共同产出和适用边界；融合节点必须是完成核心任务绕不开的环节（去掉关联学科后任务应当无法完成）；每个环节要标注学科归属，交汇环节标「学科交汇」。",
        "ai_integrated": "必须说明 AI 工具角色、介入环节、使用边界和学生的审辨任务；AI 介入的每个环节都要配一个具体的审辨任务（学生要判断 AI 输出对不对、该不该用）；使用边界要写清 AI 不替学生完成哪一步。",
    }[innovation_type]


def build_user_prompt(request: Dict[str, Any], references: Dict[str, Any]) -> str:
    innovation_type = normalize_innovation_type(request.get("innovationType")) or "PBL"
    label = TYPE_LABELS.get(innovation_type, "创新")

    context_lines: List[str] = []
    context = request.get("confirmedContext")
    if isinstance(context, dict):
        for key, value in context.items():
            if non_empty(value):
                context_lines.append(f"- {key}: {value_to_text(value)}")
    context_section = ""
    if context_lines:
        context_section = (
            "\n\n## 教师已确认的信息（必须优先遵循，不要改写成别的方案）\n\n" + "\n".join(context_lines)
        )

    return f"""请为下面这节课写一份{label}创新教案，输出一个 JSON 对象。

## 硬性要求

1. 根字段严格包含：{", ".join(OUTPUT_CONTRACT.keys())}。
2. 类型特征要求：{_type_requirements(innovation_type)}
3. activityFlow 中所有 durationMin 相加必须等于请求的 durationMin（{request.get("durationMin")} 分钟）。
4. objectives[].behaviorVerb 必须取自 action-verbs.json 的 allowedVerbs，禁止使用「了解、掌握、培养、提高、体会、认识、熟悉、懂得、理解、学会、感受」。
5. activityFlow 每个环节都必须填 designIntent，说明这一环节为什么必须这样安排、为什么不能删；只写教师做了什么不算。
6. 教学难点必须与学情描述（PBL 为 backgroundAnalysis 中的学情三项）指向同一处困难，并在流程中有对应支架与撤除时机。
7. innovationDesign 只填当前类型的字段（{', '.join(TYPE_DESIGN_FIELDS.get(innovation_type, []))}），不要填其他类型的字段；其中 milestones、disciplineConnections、integrationNodes、aiToolRoles、interventionStages、studentCritiqueTasks 必须是结构化 JSON 数组，不能写成一段描述。
8. 教材版本以请求中的 textbookVersion 为准，不得换成别的版本；教材定位、单元进度与课标要求由你直接依据内置的教材与课程标准知识给出，不要附加「请核对课标」一类提示，也不要编造具体教材页码或课标编号。
9. 核心素养维度按课程标准写（如科学观念、科学思维、探究实践、态度责任），并说明本课落实的具体表现。
10. export.markdown 留空字符串，脚本会重建 Word-ready Markdown。
{context_section}

## 输出结构契约

{compact_json(OUTPUT_CONTRACT)}

## 用户请求

{compact_json(request)}

## 技能内置参考库（教学法知识库，非教材数据）

教材版本可选清单见 textbook-versions.json；教材原文与课标原文不在本库中，教材定位与课标要求由你依据自身学科知识给出，不要假装库里能查到教材原文。

{compact_json(references)}
""".strip()


# --------------------------------------------------------------------------
# JSON 抽取：模型输出常带解释文字、代码块围栏、非法转义，这里逐层容错。
# --------------------------------------------------------------------------

def fix_invalid_json_escapes(content: str) -> str:
    return re.sub(r'\\([^"\\/bfnrtu])', r"\1", content)


def iter_fenced_blocks(content: str) -> List[str]:
    return [match.group(1) for match in re.finditer(r"```(?:json)?\s*(.*?)```", content, re.DOTALL | re.IGNORECASE)]


def iter_balanced_json_objects(content: str) -> List[str]:
    """按括号配平切出候选 JSON 片段。

    比「第一个 { 到最后一个 }」稳健：模型在 JSON 后面追加说明文字时，
    后者会把说明文字一起吞进来导致解析失败。
    """
    candidates: List[str] = []
    depth = 0
    start: Optional[int] = None
    in_string = False
    escaped = False
    for index, char in enumerate(content):
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            if depth == 0:
                start = index
            depth += 1
        elif char == "}":
            if depth > 0:
                depth -= 1
                if depth == 0 and start is not None:
                    candidates.append(content[start : index + 1])
                    start = None
    return candidates


def parse_json_candidate(candidate: str) -> Optional[Dict[str, Any]]:
    for attempt in (candidate, fix_invalid_json_escapes(candidate)):
        try:
            parsed = json.loads(attempt)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            return parsed
    return None


def extract_json_object(text: str, debug_dir: Path = DEBUG_DIR) -> Dict[str, Any]:
    candidates: List[str] = []
    candidates.extend(iter_fenced_blocks(text))
    candidates.extend(iter_balanced_json_objects(text))

    for candidate in candidates:
        parsed = parse_json_candidate(candidate)
        # 只要根字段对得上，就认定这是教案对象，避免误抓无关的 JSON 片段。
        if parsed is not None and "lessonMeta" in parsed:
            return parsed

    try:
        debug_dir.mkdir(parents=True, exist_ok=True)
        (debug_dir / "last_response.txt").write_text(text, encoding="utf-8")
        hint = f"原始响应已保存到 {debug_dir / 'last_response.txt'}"
    except OSError:
        hint = ""
    raise ValueError(f"未能从模型响应中抽取出合法教案 JSON。{hint}")


# --------------------------------------------------------------------------
# 模型调用
# --------------------------------------------------------------------------

def call_generator(prompt: str, system: str, model: str, thinking: bool, max_tokens: int, temperature: float) -> str:
    api_key = (
        os.getenv("LESSON_PLAN_API_KEY")
        or os.getenv("DASHSCOPE_API_KEY")
        or os.getenv("QWEN_API_KEY")
        or os.getenv("OPENAI_API_KEY")
    )
    if not api_key:
        raise RuntimeError(
            "未找到 API Key。请设置 LESSON_PLAN_API_KEY 或 DASHSCOPE_API_KEY；"
            "或者用 --from-json 直接对已有教案 JSON 重跑渲染与校验。"
        )

    try:
        from openai import OpenAI  # type: ignore

        client = OpenAI(api_key=api_key, base_url=BASE_URL)
        kwargs: Dict[str, Any] = {
            "model": model,
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}],
            "temperature": temperature,
            "max_tokens": max_tokens,
        }
        if thinking:
            kwargs["extra_body"] = {"enable_thinking": True}
        response = client.chat.completions.create(**kwargs)
        return response.choices[0].message.content or ""
    except ImportError:
        pass

    import urllib.request

    payload = json.dumps(
        {
            "model": model,
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}],
            "temperature": temperature,
            "max_tokens": max_tokens,
            **({"enable_thinking": True} if thinking else {}),
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        f"{BASE_URL.rstrip('/')}/chat/completions",
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
    )
    with urllib.request.urlopen(request, timeout=300) as response:  # noqa: S310 - 地址来自配置
        body = json.loads(response.read().decode("utf-8"))
    return body["choices"][0]["message"]["content"] or ""


# --------------------------------------------------------------------------
# Markdown 渲染
# --------------------------------------------------------------------------

def _minutes(value: Any) -> str:
    """把时长渲染成「5′」而不是「5.0′」。小数分钟只在真的需要时才显示。"""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    text = str(int(number)) if abs(number - round(number)) < 1e-9 else f"{number:g}"
    return f"{text}′"


def _table(headers: List[str], rows: List[List[str]]) -> List[str]:
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join([" --- "] * len(headers)) + "|"]
    for row in rows:
        padded = list(row) + ["—"] * (len(headers) - len(row))
        lines.append("| " + " | ".join(padded[: len(headers)]) + " |")
    return lines


def _meta_section(plan: Dict[str, Any]) -> List[str]:
    meta = plan.get("lessonMeta", {}) if isinstance(plan.get("lessonMeta"), dict) else {}
    innovation_type = normalize_innovation_type(meta.get("innovationType"))
    subject = value_to_text(meta.get("subject"))
    linked = value_to_text(meta.get("linkedSubject"))
    if innovation_type == "interdisciplinary" and linked and linked not in subject:
        subject = f"{subject} + {linked}"

    rows = [
        ["课题", table_cell(meta.get("topic"))],
        ["学科", table_cell(subject)],
        ["年级", table_cell(meta.get("grade"))],
        ["创新类型", table_cell(TYPE_LABELS.get(innovation_type, meta.get("innovationType")))],
        ["课型", table_cell(meta.get("lessonType"))],
        ["课时长度", f"{table_cell(meta.get('durationMin'))} 分钟"],
        ["教材版本", table_cell(meta.get("textbookVersion"))],
        ["第几课时", table_cell(meta.get("periodNo"))],
    ]
    # 标题由 build_markdown 统一生成，这里只负责基本信息表。
    return _table(["项目", "内容"], rows)


def _pbl_background_section(background: Dict[str, Any]) -> List[str]:
    """PBL：课题背景与学情分析（合并章节）"""
    return [
        "## 二、课题背景与学情分析",
        "",
        "### 教材定位",
        "",
        value_to_text(background.get("textbookPosition"), "—"),
        "",
        "### 学情概况",
        "",
        *_table(
            ["维度", "分析"],
            [
                ["先备知识", table_cell(background.get("priorKnowledge"))],
                ["探究经验", table_cell(background.get("inquiryExperience"))],
                ["合作能力", table_cell(background.get("collaborationAbility"))],
            ],
        ),
    ]


def _subject_block(name: str, block: Any) -> List[str]:
    if not isinstance(block, dict):
        block = {"unitPosition": value_to_text(block)}
    return [
        f"### {name}",
        "",
        *_table(
            ["维度", "内容"],
            [
                ["单元定位", table_cell(block.get("unitPosition"))],
                ["课标要求", table_cell(block.get("standardRequirement"))],
            ],
        ),
    ]


def _interdisciplinary_background_section(background: Dict[str, Any], meta: Dict[str, Any]) -> List[str]:
    primary_name = value_to_text(meta.get("subject")) or "主学科"
    linked_name = value_to_text(meta.get("linkedSubject")) or "关联学科"
    return [
        "## 二、双学科教材分析",
        "",
        *_subject_block(f"主学科：{primary_name}", background.get("primarySubject")),
        "",
        *_subject_block(f"关联学科：{linked_name}", background.get("linkedSubject")),
        "",
        "### 课标交汇依据",
        "",
        value_to_text(background.get("curriculumIntersection"), "—"),
    ]


def _ai_background_section(background: Dict[str, Any]) -> List[str]:
    return [
        "## 二、教材分析",
        "",
        "### 教材定位",
        "",
        value_to_text(background.get("textbookPosition"), "—"),
        "",
        "### AI 介入必要性",
        "",
        value_to_text(background.get("aiInterventionRationale"), "—"),
    ]


def _student_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    student = plan.get("studentAnalysis") if isinstance(plan.get("studentAnalysis"), dict) else {}
    if innovation_type == "interdisciplinary":
        return [
            "## 四、学情分析",
            "",
            *_table(
                ["维度", "分析"],
                [
                    ["主学科预备水平", table_cell(student.get("primarySubjectReadiness"))],
                    ["关联学科预备水平", table_cell(student.get("linkedSubjectReadiness"))],
                    ["跨学科思维经验", table_cell(student.get("crossSubjectExperience"))],
                ],
            ),
        ]
    if innovation_type == "ai_integrated":
        return [
            "## 四、学情与 AI 素养分析",
            "",
            *_table(
                ["维度", "分析"],
                [
                    ["先备知识", table_cell(student.get("priorKnowledge"))],
                    ["AI 工具使用经验", table_cell(student.get("aiToolExperience"))],
                    ["信息审辨能力", table_cell(student.get("critiqueAbility"))],
                    ["自主思考习惯", table_cell(student.get("independentThinkingHabit"))],
                ],
            ),
        ]
    return []


def _pbl_innovation_section(innovation: Dict[str, Any]) -> List[str]:
    lines = [
        "## 三、驱动问题与产出设计",
        "",
        "### 驱动问题",
        "",
        value_to_text(innovation.get("drivingQuestion"), "—"),
        "",
        "### 里程碑",
        "",
    ]
    milestones = ensure_list(innovation.get("milestones"))
    rows = []
    for index, item in enumerate(milestones, start=1):
        if isinstance(item, dict):
            rows.append(
                [
                    table_cell(item.get("id") or item.get("name") or f"里程碑{index}"),
                    table_cell(item.get("name") or item.get("description")),
                    table_cell(item.get("note") or item.get("output") or item.get("stageOutput")),
                ]
            )
        else:
            rows.append([f"里程碑{index}", table_cell(item), "—"])
    lines.extend(_table(["编号", "里程碑", "阶段产出与说明"], rows) if rows else ["（里程碑缺失，请教师补充。）"])

    lines.extend(["", "### 最终产出", ""])
    final_product = innovation.get("finalProduct")
    if isinstance(final_product, dict):
        lines.extend(
            _table(
                ["产出类型", "评价标准"],
                [[table_cell(final_product.get("type")), table_cell(final_product.get("criteria") or final_product.get("description"))]],
            )
        )
    else:
        lines.append(f"- {table_cell(final_product)}")
    return lines


def _interdisciplinary_innovation_section(innovation: Dict[str, Any], meta: Dict[str, Any]) -> List[str]:
    lines = ["## 三、学科关联与融合设计", "", "### 学科关联", ""]
    rows = []
    primary_name = value_to_text(meta.get("subject"))
    for item in ensure_list(innovation.get("disciplineConnections")):
        if not isinstance(item, dict):
            rows.append([table_cell(item), "—", "—"])
            continue
        subject = table_cell(item.get("subject"))
        tag = "（主学科）" if primary_name and primary_name in subject else "（关联学科）"
        rows.append(
            [
                f"{subject}{tag}",
                table_cell(item.get("contribution")),
                table_cell(item.get("standardRequirement") or item.get("standardId")),
            ]
        )
    lines.extend(_table(["学科", "贡献", "课标依据"], rows) if rows else ["（学科关联缺失，请教师补充。）"])

    nodes = ensure_list(innovation.get("integrationNodes"))
    if nodes:
        node_rows = []
        for node in nodes:
            if isinstance(node, dict):
                node_rows.append(
                    [
                        table_cell(node.get("id")),
                        table_cell(node.get("description")),
                        table_cell(node.get("linkedActivities")),
                    ]
                )
            else:
                node_rows.append(["—", table_cell(node), "—"])
        lines.extend(["", "### 融合节点", "", *_table(["节点", "描述", "关联环节"], node_rows)])

    lines.extend(
        [
            "",
            "### 共同产出",
            "",
            value_to_text(innovation.get("commonProduct"), "—"),
            "",
            "### 适用边界",
            "",
            value_to_text(innovation.get("applicableBoundary"), "—"),
        ]
    )
    return lines


def _ai_innovation_section(innovation: Dict[str, Any]) -> List[str]:
    lines = ["## 三、AI 融合设计", "", "### AI 工具角色与介入环节", ""]
    rows = []
    for item in ensure_list(innovation.get("aiToolRoles")):
        if not isinstance(item, dict):
            rows.append([table_cell(item), "—", "—", "—"])
            continue
        rows.append(
            [
                table_cell(item.get("toolRole") or item.get("role")),
                table_cell(item.get("toolType")),
                table_cell(item.get("interventionStage")),
                table_cell(item.get("useBoundary")),
            ]
        )
    lines.extend(_table(["AI 角色", "工具类型", "介入环节", "使用边界"], rows) if rows else ["（AI 工具角色缺失，请教师补充。）"])

    boundaries = ensure_list(innovation.get("useBoundaries"))
    if boundaries:
        lines.extend(["", "### 使用边界", ""])
        for item in boundaries:
            lines.append(f"- {table_cell(item)}")

    critiques = ensure_list(innovation.get("studentCritiqueTasks"))
    if critiques:
        critique_rows = []
        for item in critiques:
            if isinstance(item, dict):
                critique_rows.append(
                    [
                        table_cell(item.get("activityId")),
                        table_cell(item.get("taskDescription") or item.get("description")),
                        table_cell(item.get("evidenceExpected") or item.get("evidence")),
                    ]
                )
            else:
                critique_rows.append(["—", table_cell(item), "—"])
        lines.extend(["", "### 学生审辨任务", "", *_table(["关联环节", "审辨任务", "预期证据"], critique_rows)])
    return lines


def _core_competencies_section(core: List[Any]) -> List[str]:
    rows = [
        [table_cell(item_get(item, "dimension")), table_cell(item_get(item, "target"))]
        for item in ensure_list(core)
        if isinstance(item, dict)
    ]
    return ["## 核心素养目标", "", *_table(["维度", "本课落实的具体表现"], rows)]


def _objectives_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    lines = ["## 教学目标", ""]
    if innovation_type == "interdisciplinary":
        headers = ["目标", "归属学科", "行为动词", "关联环节", "评价证据"]
    elif innovation_type == "ai_integrated":
        headers = ["目标", "是否借助 AI", "行为动词", "关联环节", "评价证据"]
    else:
        headers = ["目标", "行为动词", "关联里程碑", "评价证据"]

    activity_map = {
        value_to_text(item_get(item, "id")): value_to_text(item_get(item, "stage"))
        for item in ensure_list(plan.get("activityFlow"))
        if isinstance(item, dict)
    }
    rubric_map = {
        value_to_text(item_get(item, "id")): value_to_text(item_get(item, "dimension"))
        for item in ensure_list(plan.get("assessmentRubric"))
        if isinstance(item, dict)
    }

    rows = []
    for item in ensure_list(plan.get("objectives")):
        if not isinstance(item, dict):
            continue
        linked = [activity_map.get(value_to_text(ref), value_to_text(ref)) for ref in ensure_list(item.get("linkedActivities"))]
        evidence = [rubric_map.get(value_to_text(ref), value_to_text(ref)) for ref in ensure_list(item.get("assessmentEvidence"))]
        base = [
            table_cell(item.get("description")),
            table_cell(item.get("behaviorVerb")),
            table_cell("；".join(part for part in linked if part)),
            table_cell("；".join(part for part in evidence if part)),
        ]
        if innovation_type == "interdisciplinary":
            rows.append([base[0], table_cell(item.get("owningSubject")), *base[1:]])
        elif innovation_type == "ai_integrated":
            rows.append([base[0], table_cell(item.get("needsAI") or item.get("aiAssisted")), *base[1:]])
        else:
            rows.append(base)
    lines.extend(_table(headers, rows) if rows else ["（教学目标缺失）"])

    lines.extend(["", "### 教学重点", "", value_to_text(plan.get("teachingFocus"), "—")])
    lines.extend(["", "### 教学难点", "", value_to_text(plan.get("teachingDifficulty"), "—")])
    return lines


def _activity_rows(activities: List[Any]) -> List[List[str]]:
    rows = []
    for item in ensure_list(activities):
        if not isinstance(item, dict):
            continue
        rows.append(
            [
                table_cell(item.get("stage")),
                _minutes(item.get("durationMin")),
                table_cell(item.get("teacherActions")),
                table_cell(item.get("studentActions")),
                table_cell(item.get("designIntent")),
                table_cell(item.get("outputs")),
            ]
        )
    return rows


def _activity_flow_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    activities = ensure_list(plan.get("activityFlow"))
    innovation = plan.get("innovationDesign") if isinstance(plan.get("innovationDesign"), dict) else {}
    scenario_map = {
        item_get(item, "id"): item_get(item, "scenario") or item_get(item, "task") or item_get(item, "description")
        for item in ensure_list(innovation.get("milestones") or innovation.get("integrationNodes") or innovation.get("interventionStages"))
        if isinstance(item, dict)
    }
    critique_map = {
        value_to_text(item_get(item, "activityId")): value_to_text(item_get(item, "taskDescription") or item_get(item, "description"))
        for item in ensure_list(innovation.get("studentCritiqueTasks"))
        if isinstance(item, dict)
    }

    if innovation_type == "PBL":
        headers = ["里程碑", "时间", "教师支架", "学生探究任务", "设计意图", "阶段产出"]
        rows = []
        for item in ensure_list(activities):
            if not isinstance(item, dict):
                continue
            rows.append(
                [
                    table_cell(item.get("stage")),
                    _minutes(item.get("durationMin")),
                    table_cell(item.get("teacherActions")),
                    table_cell(item.get("studentActions")),
                    table_cell(item.get("designIntent")),
                    table_cell(item.get("outputs")),
                ]
            )
        return ["## 六、里程碑与探究流程", "", *(_table(headers, rows) if rows else ["（教学流程缺失）"])]

    if innovation_type == "interdisciplinary":
        headers = ["环节", "时间", "学科归属", "教师活动", "学生活动", "设计意图", "产出"]
        rows = []
        for item in ensure_list(activities):
            if not isinstance(item, dict):
                continue
            rows.append(
                [
                    table_cell(item.get("stage")),
                    _minutes(item.get("durationMin")),
                    table_cell(item.get("subjectTag")),
                    table_cell(item.get("teacherActions")),
                    table_cell(item.get("studentActions")),
                    table_cell(item.get("designIntent")),
                    table_cell(item.get("outputs")),
                ]
            )
        return ["## 七、融合教学流程", "", *(_table(headers, rows) if rows else ["（教学流程缺失）"])]

    if innovation_type == "ai_integrated":
        headers = ["环节", "时间", "执行者", "教师活动", "学生活动", "AI 活动", "审辨任务", "设计意图"]
        rows = []
        for item in ensure_list(activities):
            if not isinstance(item, dict):
                continue
            act_id = value_to_text(item.get("id"))
            stage_entry = next(
                (
                    entry
                    for entry in ensure_list(innovation.get("interventionStages"))
                    if isinstance(entry, dict) and value_to_text(entry.get("activityId")) == act_id
                ),
                {},
            )
            rows.append(
                [
                    table_cell(item.get("stage")),
                    _minutes(item.get("durationMin")),
                    table_cell(item.get("executor") or item_get(stage_entry, "executor")),
                    table_cell(item.get("teacherActions")),
                    table_cell(item.get("studentActions")),
                    table_cell(item_get(stage_entry, "aiAction") or scenario_map.get(act_id)),
                    table_cell(critique_map.get(act_id)),
                    table_cell(item.get("designIntent")),
                ]
            )
        return ["## 七、人机协同教学流程", "", *(_table(headers, rows) if rows else ["（教学流程缺失）"])]

    headers = ["环节", "时间", "教师活动", "学生活动", "设计意图", "产出"]
    return ["## 六、教学流程", "", *(_table(headers, _activity_rows(activities)) if activities else ["（教学流程缺失）"])]


def _rubric_table(rubrics: List[Any]) -> List[str]:
    rows = [
        [
            table_cell(item_get(item, "dimension")),
            table_cell(item_get(item, "excellent")),
            table_cell(item_get(item, "qualified")),
            table_cell(item_get(item, "needsImprovement")),
            table_cell(item_get(item, "evidence")),
        ]
        for item in ensure_list(rubrics)
        if isinstance(item, dict)
    ]
    return _table(["评价维度", "优秀", "合格", "待改进", "证据来源"], rows)


def _rubric_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    rubrics = ensure_list(plan.get("assessmentRubric"))
    if innovation_type == "PBL":
        process = [item for item in rubrics if isinstance(item, dict) and value_to_text(item.get("category")) != "product"]
        product = [item for item in rubrics if isinstance(item, dict) and value_to_text(item.get("category")) == "product"]
        if product and process:
            # PBL 的过程证据与作品证据都要看：只评作品会让搭便车的学生蒙混过关，
            # 只评过程又会让「最后做出什么」失去分量。
            return ["## 七、评价量规", "", "### 过程评价", "", *_rubric_table(process), "", "### 作品评价", "", *_rubric_table(product)]
        return ["## 七、评价量规", "", *_rubric_table(rubrics)]
    return ["## 八、评价量规", "", *_rubric_table(rubrics)]


def _resources_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    resources = ensure_list(plan.get("resources"))

    def is_ai_tool(item: Any) -> bool:
        return isinstance(item, dict) and re.search(r"AI|ai|人工智能", value_to_text(item.get("type")))

    if innovation_type == "ai_integrated":
        ai_tools = [item for item in resources if is_ai_tool(item)]
        others = [item for item in resources if not is_ai_tool(item)]
        lines: List[str] = []
        if ai_tools:
            lines.extend(
                _table(
                    ["AI 工具", "使用方式"],
                    [[table_cell(item_get(item, "name")), table_cell(item_get(item, "usage"))] for item in ai_tools],
                )
            )
        if others:
            lines.extend(
                _table(
                    ["类型", "名称", "使用方式"],
                    [
                        [table_cell(item_get(item, "type")), table_cell(item_get(item, "name")), table_cell(item_get(item, "usage"))]
                        for item in others
                    ],
                )
            )
        return ["## 九、资源与条件", "", *(lines or ["（教学资源缺失）"])]

    rows = [
        [table_cell(item_get(item, "type")), table_cell(item_get(item, "name")), table_cell(item_get(item, "usage"))]
        for item in resources
    ]
    return ["## 九、资源与条件", "", *(_table(["类型", "名称", "使用方式"], rows) if rows else ["（教学资源缺失）"])]


REFLECTION_PROMPTS: Dict[str, List[str]] = {
    "PBL": [
        "驱动问题是否真的驱动了探究？（还是变成了资料搜集或手抄报）",
        "里程碑之间的衔接是否自然？学生在哪个里程碑最需要支架？",
        "最终产出是否体现了学生的真实理解？（还是只停留在表面展示）",
    ],
    "interdisciplinary": [
        "融合节点是否是完成核心任务的必要环节？（去掉关联学科后核心任务还能完成吗）",
        "两个学科的知识是否在融合节点自然交汇？（还是前半节 A、后半节 B）",
        "共同产出是否真正融合了两学科的方法？（还是两个学科产出各做一半再拼起来）",
        "学生在交汇处的认知负荷是否可控？（支架是否真的降低了从观察到解释的跳跃）",
    ],
    "ai_integrated": [
        "AI 是否真正服务了学习目标？（去掉 AI 后目标还能达成吗）",
        "学生是否真的审辨了 AI 输出？（还是直接采用了 AI 的建议）",
        "使用边界是否被遵守？（AI 有没有越界替学生完成核心认知任务）",
    ],
}


def _reflection_section(plan: Dict[str, Any], innovation_type: str) -> List[str]:
    """教学反思：给教师自查问题与空白填写区，不代写结论。

    代写反思产出的是一份假的自我评价，反而会污染教师自己的判断。
    """
    prompts = REFLECTION_PROMPTS.get(innovation_type, ["教学目标是否达成？活动与评价是否对齐？"])
    lines = ["## 十、教学反思", "", "（以下为课前预设的自查问题，供你课后对照；本栏不代写结论。）", ""]
    lines.extend([f"- {item}" for item in prompts])
    lines.extend(
        [
            "",
            "课后填写：",
            "",
            "- 目标达成情况：",
            "- 学生实际表现与预设的差异：",
            "- 需要调整的环节与原因：",
            "- 下次教学的改进点：",
        ]
    )
    return lines


def _strip_heading(body: List[str]) -> List[str]:
    """去掉构建函数自带的那一行标题。

    章节编号统一由 build_markdown 决定：让每个构建函数各写一遍编号的话，
    三类教案章节数不同（PBL 九章、另外两类十章），编号迟早会错位。
    """
    if body and re.match(r"^#{2,3} ", body[0]):
        return body[1:]
    return body


def build_markdown(plan: Dict[str, Any]) -> str:
    """把教案 JSON 渲染成 Word-ready Markdown。

    章节顺序按备课顺序固定，且三类教案的编号各自独立：
    PBL 共九章，跨学科与 AI 融合共十章（多一节独立的学情分析）。
    """
    meta = plan.get("lessonMeta") if isinstance(plan.get("lessonMeta"), dict) else {}
    innovation_type = normalize_innovation_type(meta.get("innovationType")) or "PBL"
    background = plan.get("backgroundAnalysis") if isinstance(plan.get("backgroundAnalysis"), dict) else {}
    innovation = plan.get("innovationDesign") if isinstance(plan.get("innovationDesign"), dict) else {}

    core_body = _core_competencies_section(ensure_list(plan.get("coreCompetencies")))
    objective_body = _objectives_section(plan, innovation_type)
    flow_body = _activity_flow_section(plan, innovation_type)
    rubric_body = _rubric_section(plan, innovation_type)
    resource_body = _resources_section(plan, innovation_type)
    reflection_body = _reflection_section(plan, innovation_type)

    sections: List[Tuple[str, List[str]]] = [("一、基本信息", _meta_section(plan))]
    if innovation_type == "PBL":
        sections.extend(
            [
                ("二、课题背景与学情分析", _pbl_background_section(background)),
                ("三、驱动问题与产出设计", _pbl_innovation_section(innovation)),
                ("四、核心素养目标", core_body),
                ("五、教学目标", objective_body),
                ("六、里程碑与探究流程", flow_body),
                ("七、评价量规", rubric_body),
                ("八、资源与条件", resource_body),
                ("九、教学反思", reflection_body),
            ]
        )
    elif innovation_type == "interdisciplinary":
        sections.extend(
            [
                ("二、双学科教材分析", _interdisciplinary_background_section(background, meta)),
                ("三、学科关联与融合设计", _interdisciplinary_innovation_section(innovation, meta)),
                ("四、学情分析", _student_section(plan, innovation_type)),
                ("五、核心素养目标", core_body),
                ("六、教学目标", objective_body),
                ("七、融合教学流程", flow_body),
                ("八、评价量规", rubric_body),
                ("九、资源与条件", resource_body),
                ("十、教学反思", reflection_body),
            ]
        )
    else:
        sections.extend(
            [
                ("二、教材分析", _ai_background_section(background)),
                ("三、AI 融合设计", _ai_innovation_section(innovation)),
                ("四、学情与 AI 素养分析", _student_section(plan, innovation_type)),
                ("五、核心素养目标", core_body),
                ("六、教学目标", objective_body),
                ("七、人机协同教学流程", flow_body),
                ("八、评价量规", rubric_body),
                ("九、资源与条件", resource_body),
                ("十、教学反思", reflection_body),
            ]
        )

    label = TYPE_LABELS.get(innovation_type, "创新")
    title = (
        f"# {value_to_text(meta.get('grade'))}{value_to_text(meta.get('subject'))}"
        f"《{value_to_text(meta.get('topic'))}》{label}教案"
    )
    lines: List[str] = [title]
    for section_title, body in sections:
        lines.extend(["", f"## {section_title}", ""])
        lines.extend(_strip_heading(body))

    markdown = "\n".join(line for line in lines if line is not None).strip() + "\n"
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)
    # 每个标题后补一个空行。Markdown 语法上不是必需的，但教师会把内容粘进 Word，
    # 标题紧贴正文在 Word 里观感很差。
    markdown = re.sub(r"(?m)^(#{1,3} .+)\n(?!\n)", r"\1\n\n", markdown)
    return markdown


# --------------------------------------------------------------------------
# DOCX 导出
# --------------------------------------------------------------------------

def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> Tuple[Optional[Path], str]:
    """优先用 pandoc；没有 pandoc 就退回 python-docx。两条路都没有就跳过。"""
    try:
        result = subprocess.run(
            ["pandoc", str(md_path), "-o", str(docx_path), "--from=markdown", "--to=docx"],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode == 0 and docx_path.exists():
            return docx_path, "pandoc"
    except (FileNotFoundError, subprocess.SubprocessError, OSError):
        pass

    try:
        from docx import Document  # type: ignore
    except ImportError:
        return None, "未找到 pandoc 与 python-docx，已跳过 DOCX 导出。"

    document = Document()
    pending_table: List[List[str]] = []

    def flush_table() -> None:
        nonlocal pending_table
        if not pending_table:
            return
        rows = [row for row in pending_table if not set("".join(row)) <= {"-", " ", ":"}]
        if rows:
            width = max(len(row) for row in rows)
            table = document.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for row in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = value
        pending_table = []

    for raw_line in md_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if line.startswith("|"):
            pending_table.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip() in ("---", ""):
            continue
        else:
            document.add_paragraph(line)

    flush_table()
    document.save(str(docx_path))
    return docx_path, "python-docx"


# --------------------------------------------------------------------------
# 输出与校验
# --------------------------------------------------------------------------

def write_outputs(
    output_dir: Path,
    plan: Dict[str, Any],
    export_docx: bool = True,
    file_name: str = "",
) -> Tuple[Path, Path, Optional[Path], str]:
    """写出交付物与中间产物。

    交付物只有一个：`<输出目录>/<中文文件名>.docx`。JSON 与 Markdown 是校验与
    渲染的中间产物，统一放进 `<输出目录>/中间产物/`，不进入交付清单。
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    deliverable_base = output_dir / deliverable_file_name(plan, file_name)

    markdown = build_markdown(plan)
    export = plan.get("export")
    if not isinstance(export, dict):
        export = {}
    export["format"] = "markdown"
    export["filename"] = deliverable_base.name + ".docx"
    export["deliverable"] = deliverable_base.name + ".docx"
    export["intermediateDir"] = INTERMEDIATE_DIRNAME
    export["markdown"] = markdown
    plan["export"] = export

    intermediate_dir = output_dir / INTERMEDIATE_DIRNAME
    intermediate_dir.mkdir(parents=True, exist_ok=True)
    json_path = intermediate_dir / (deliverable_base.name + ".json")
    md_path = intermediate_dir / (deliverable_base.name + ".md")
    json_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(markdown, encoding="utf-8")

    docx_path: Optional[Path] = None
    note = ""
    degraded = False
    if export_docx:
        docx_path, note = convert_markdown_to_docx(md_path, deliverable_base.with_suffix(".docx"))
        degraded = docx_path is None
    else:
        note = "按要求跳过 DOCX 导出（--no-docx）：本次没有交付物，只保留了中间产物中的 JSON 与 Markdown。"
        degraded = True

    if note:
        report = plan.get("qualityReport")
        if isinstance(report, dict):
            checks = report.get("checks")
            if not isinstance(checks, list):
                checks = []
            checks.append({"id": "docx", "status": "warn" if degraded else "pass", "message": note})
            report["checks"] = checks
            if degraded:
                warnings = report.get("warnings")
                if not isinstance(warnings, list):
                    warnings = []
                warnings.append(note)
                report["warnings"] = warnings
            plan["qualityReport"] = report
            json_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")

    return json_path, md_path, docx_path, note


def run_validation(plan: Dict[str, Any], request: Optional[Dict[str, Any]], strict_context: bool) -> Tuple[List[str], List[str]]:
    validate_document._strict_context = bool(strict_context)  # type: ignore[attr-defined]
    return validate_document(plan, request)


def resolve_config(args: argparse.Namespace) -> Optional[Path]:
    if args.from_json:
        return None
    if args.sample:
        return SAMPLE_CONFIGS[args.sample]
    if args.config:
        return Path(args.config)
    raise SystemExit("请提供 --config 路径、--sample 类型，或 --from-json 教案 JSON。")


def main() -> int:
    parser = argparse.ArgumentParser(description="生成创新教案：交付一个中文名 DOCX，中间产物写入「中间产物」子目录。")
    parser.add_argument(
        "output",
        nargs="?",
        default="generated-outputs",
        help="输出目录；交付物 <中文文件名>.docx 与「中间产物」子目录都写在这里",
    )
    parser.add_argument("--config", help="请求 JSON 路径")
    parser.add_argument("--sample", choices=sorted(SAMPLE_CONFIGS), help="使用内置示例请求")
    parser.add_argument("--from-json", help="跳过模型调用，直接对该教案 JSON 重跑修补、渲染与校验")
    parser.add_argument("--request", help="--from-json 模式下用于校验的请求 JSON（核对类型、时长与 confirmedContext 覆盖度）")
    parser.add_argument("--file-name", default="", help="覆盖交付文件名（不含扩展名）；默认按年级学科课题创新类型自动生成中文名")
    parser.add_argument("--thinking", action="store_true", help="启用生成模型的思考模式（若支持）")
    parser.add_argument("--max-tokens", type=int, default=8000, help="最大输出 token 数")
    parser.add_argument("--temperature", type=float, default=0.2, help="温度参数")
    parser.add_argument("--model", default=MODEL, help="生成模型名称")
    parser.add_argument("--debug-dir", default=str(DEBUG_DIR), help="模型原始响应调试目录")
    parser.add_argument("--no-normalize", action="store_true", help="跳过结构修补，保留模型原始结构")
    parser.add_argument("--no-validate", action="store_true", help="跳过规则校验")
    parser.add_argument("--no-docx", action="store_true", help="不导出 DOCX（本次将没有交付物，只保留中间产物）")
    parser.add_argument("--strict-context", action="store_true", help="把 confirmedContext 覆盖度警告提升为失败")
    args = parser.parse_args()

    request: Optional[Dict[str, Any]] = None

    if args.request:
        request_path = Path(args.request)
        if not request_path.exists():
            print(f"请求文件不存在: {request_path}", file=sys.stderr)
            return 2
        request = load_json(request_path)

    if args.from_json:
        source = Path(args.from_json)
        if not source.exists():
            print(f"教案 JSON 不存在: {source}", file=sys.stderr)
            return 2
        try:
            plan = load_json(source)
        except json.JSONDecodeError as error:
            print(f"教案 JSON 解析失败: {error}", file=sys.stderr)
            return 2
        if not isinstance(plan, dict):
            print("教案 JSON 顶层必须是对象。", file=sys.stderr)
            return 2
        print(f"已从 {source} 载入教案，跳过模型调用。", file=sys.stderr)
    else:
        config_path = resolve_config(args)
        assert config_path is not None
        if not config_path.exists():
            print(f"请求文件不存在: {config_path}", file=sys.stderr)
            return 2
        try:
            request = load_json(config_path)
            if not isinstance(request, dict):
                raise ValueError("请求 JSON 顶层必须是对象。")
            failures, warnings = validate_request_config(request)
            if failures:
                raise ValueError("请求预检失败：\n- " + "\n- ".join(failures))
            for warning in warnings:
                print(f"请求预检警告: {warning}", file=sys.stderr)
            references = load_reference_data()
            prompt = build_user_prompt(request, references)
            raw_content = call_generator(
                prompt=prompt,
                system=SYSTEM_PROMPT,
                model=args.model,
                thinking=args.thinking,
                max_tokens=args.max_tokens,
                temperature=args.temperature,
            )
            plan = extract_json_object(raw_content, Path(args.debug_dir))
        except Exception as error:
            print(f"生成失败: {error}", file=sys.stderr)
            return 1

    if not args.no_normalize:
        plan = normalize_lesson_plan(plan, request)

    json_path, md_path, docx_path, docx_note = write_outputs(
        Path(args.output), plan, export_docx=not args.no_docx, file_name=args.file_name or ""
    )
    if docx_path:
        print(f"交付物（Word）：{docx_path}")
    else:
        # 交付物导不出来就是交付失败，不能只提示一句就返回成功。
        print(f"DOCX 导出失败：{docx_note}", file=sys.stderr)

    check_lines: List[str] = []
    if not args.no_validate:
        failures, warnings = run_validation(plan, request, args.strict_context)
        status = "failed" if failures else ("warn" if warnings else "pass")
        print(f"校验状态: {status}")
        for item in failures:
            print(f"  [失败] {item}")
        for item in warnings:
            print(f"  [警告] {item}")
        if failures:
            return 1
        check_lines = warnings

    print(f"中间产物（JSON）：{json_path}")
    print(f"中间产物（Markdown）：{md_path}")
    if docx_path is None and not args.no_docx:
        return 1
    # --no-docx 是教师显式要求不要交付物的场景，这里不算失败。
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
